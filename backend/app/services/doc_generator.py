import os
import logging
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from app.config import settings
from app.services.diagram_generator import (
    generate_process_flowchart,
    generate_architecture_diagram,
    generate_roadmap_timeline,
    generate_current_vs_target_diagram
)

logger = logging.getLogger("digitalization_advisor.doc_generator")

def _set_cell_background(cell, fill_hex):
    """Sets cell background color in docx table."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def _set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Sets internal padding for a cell (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

class DocumentGenerator:
    def __init__(self):
        self.exports_dir = settings.EXPORTS_DIR
        os.makedirs(self.exports_dir, exist_ok=True)

    def _add_diagram_image(self, doc, image_path: str, caption: str):
        """Embeds a generated flowchart or architecture diagram image into the document with styled caption."""
        if image_path and os.path.exists(image_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run()
            run.add_picture(image_path, width=Inches(6.0))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(14)
            r_cap = p_cap.add_run(f"Figure: {caption}")
            r_cap.font.name = "Arial"
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.bold = True
            r_cap.font.color.rgb = RGBColor(15, 118, 110)

    def _add_styled_header(self, doc, title: str, subtitle: str, doc_code: str = "SPEC-2026-001"):
        """Adds a high-end executive header banner and metadata table to the Word document."""
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(4)
        run_title = p_title.add_run(title)
        run_title.font.name = "Arial"
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 118, 110)  # Deep Emerald Teal

        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(10)
        run_sub = p_sub.add_run(f"{subtitle}  |  Ref: {doc_code}  |  Generated: {datetime.now().strftime('%B %d, %Y')}")
        run_sub.font.name = "Arial"
        run_sub.font.size = Pt(9.0)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(71, 85, 105)  # Slate gray

        # Document Control Table (Professional Emerald Header)
        table = doc.add_table(rows=2, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Document ID", "Version", "Status", "Classification"]
        values = [doc_code, "1.0 (Approved)", "Executive Specification", "Confidential & Proprietary"]
        
        for idx, text in enumerate(headers):
            cell = table.rows[0].cells[idx]
            cell.text = text
            _set_cell_background(cell, "0F766E")  # Deep Teal Emerald
            _set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

        for idx, text in enumerate(values):
            cell = table.rows[1].cells[idx]
            cell.text = text
            _set_cell_background(cell, "F0FDF4")  # Soft mint tint
            _set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = RGBColor(30, 41, 59)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def _add_green_callout_box(self, doc, title: str, text: str):
        """Adds a mint-greenish callout box for executive highlights & key metrics."""
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        _set_cell_background(cell, "ECFDF5")  # Mint green background
        _set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(0)
        
        r_head = p.add_run(f"❇️ {title}\n")
        r_head.font.name = "Arial"
        r_head.font.size = Pt(10)
        r_head.font.bold = True
        r_head.font.color.rgb = RGBColor(6, 95, 70)  # Dark Emerald

        r_text = p.add_run(text)
        r_text.font.name = "Arial"
        r_text.font.size = Pt(9.0)
        r_text.font.italic = True
        r_text.font.color.rgb = RGBColor(15, 118, 110)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def _add_kpis_green_box(self, doc, kpi_list: list):
        """Adds a dedicated Mint-Greenish Executive Key Metrics & KPIs Box into the document."""
        if not kpi_list:
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        _set_cell_background(cell, "D1FAE5")  # Rich mint green background
        _set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.2
        
        r_head = p.add_run("📊 Target Executive KPIs & Performance Metrics Benchmark\n")
        r_head.font.name = "Arial"
        r_head.font.size = Pt(10.5)
        r_head.font.bold = True
        r_head.font.color.rgb = RGBColor(6, 95, 70)

        for kpi in kpi_list:
            r_kpi = p.add_run(f"• {kpi}\n")
            r_kpi.font.name = "Arial"
            r_kpi.font.size = Pt(9.0)
            r_kpi.font.bold = True
            r_kpi.font.color.rgb = RGBColor(4, 120, 87)
            
        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def _add_benefits_roi_box(self, doc, success_criteria_benefits: dict = None, context: dict = None):
        """Adds a dedicated Key Business Benefits & Projected ROI section with styled green callout boxes."""
        self._add_heading(doc, "Key Business Benefits & Projected ROI", level=1)

        kpis = (context.get("kpis") if context else []) or (success_criteria_benefits.get("key_performance_indicators") if success_criteria_benefits else [])
        benefits = (context.get("expected_benefits") if context else []) or (success_criteria_benefits.get("business_benefits") if success_criteria_benefits else [])
        risks = (context.get("risks_and_mitigations") if context else []) or (success_criteria_benefits.get("risk_mitigations") if success_criteria_benefits else [])

        if kpis:
            self._add_kpis_green_box(doc, kpis)

        if benefits:
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            _set_cell_background(cell, "ECFDF5")
            _set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
            p = cell.paragraphs[0]
            r_head = p.add_run("📈 Strategic & Financial ROI Benefits\n")
            r_head.font.name = "Arial"
            r_head.font.size = Pt(10.5)
            r_head.font.bold = True
            r_head.font.color.rgb = RGBColor(6, 95, 70)
            for b in benefits:
                r_item = p.add_run(f"• {b}\n")
                r_item.font.name = "Arial"
                r_item.font.size = Pt(9.0)
                r_item.font.color.rgb = RGBColor(4, 120, 87)
            doc.add_paragraph().paragraph_format.space_after = Pt(8)

        if risks:
            tbl_r = doc.add_table(rows=1, cols=1)
            tbl_r.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell_r = tbl_r.rows[0].cells[0]
            _set_cell_background(cell_r, "FFFBEB")
            _set_cell_margins(cell_r, top=140, bottom=140, left=180, right=180)
            p_r = cell_r.paragraphs[0]
            r_head_r = p_r.add_run("🛡️ Risk Mitigation & Operational Controls\n")
            r_head_r.font.name = "Arial"
            r_head_r.font.size = Pt(10.5)
            r_head_r.font.bold = True
            r_head_r.font.color.rgb = RGBColor(146, 64, 14)
            for r in risks:
                r_str = f"• {r.get('risk', '')}: {r.get('mitigation', '')}" if isinstance(r, dict) else f"• {r}"
                r_item_r = p_r.add_run(f"{r_str}\n")
                r_item_r.font.name = "Arial"
                r_item_r.font.size = Pt(9.0)
                r_item_r.font.color.rgb = RGBColor(180, 83, 9)
            doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def _add_heading(self, doc, text: str, level=1):
        """Adds custom styled heading with emerald accents."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(15, 118, 110)  # Deep Emerald
        elif level == 2:
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
        else:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def generate_brd(self, session_id: str, brd_data: dict, context: dict = None) -> str:
        """Generates visual Business Requirement Document (.docx) featuring embedded flowcharts, matrices, and minimal prose."""
        doc = Document()
        doc_code = brd_data.get("doc_id_code", f"BRD-2026-{session_id[:6].upper()}")
        sol_name = context.get("solution_name") if context else brd_data.get("title", "Business Requirement Document")
        
        self._add_styled_header(doc, f"Business Requirement Document - {sol_name}", "Digitalization Advisor - Visual Solution Architecture Specification", doc_code)

        exec_sum = context.get("executive_summary") if context else brd_data.get("executive_summary", "")
        if exec_sum:
            self._add_green_callout_box(doc, "Executive Strategic Rationale", exec_sum)

        # 1. Problem Statement & Operational Friction (Visual Summary)
        self._add_heading(doc, "1. Business Problem & Operational Friction", level=1)
        prob_text = context.get("problem_statement") if context else brd_data.get("problem_statement", "")
        
        tbl_prob = doc.add_table(rows=1, cols=2)
        tbl_prob.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_left, c_right = tbl_prob.rows[0].cells
        c_left.text = "Operational Friction Point"
        c_right.text = "Target Solution Rationale"
        _set_cell_background(c_left, "0F766E")
        _set_cell_background(c_right, "0F766E")
        _set_cell_margins(c_left, top=100, bottom=100, left=120, right=120)
        _set_cell_margins(c_right, top=100, bottom=100, left=120, right=120)
        c_left.paragraphs[0].runs[0].font.bold = True
        c_left.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c_right.paragraphs[0].runs[0].font.bold = True
        c_right.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

        row_p = tbl_prob.add_row().cells
        row_p[0].text = prob_text or "High manual overhead, fragmented data sources, lack of real-time visibility."
        row_p[1].text = f"Automate end-to-end execution utilizing {sol_name} to standardize workflows and increase operational efficiency."
        for c in row_p:
            _set_cell_background(c, "F0FDF4")
            _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.size = Pt(9.0)

        # 2. Flowchart: Solution Workflow Execution
        self._add_heading(doc, "2. End-to-End Solution Workflow Architecture", level=1)
        workflow_steps = context.get("workflow_steps", []) if context else brd_data.get("workflow_steps", [])
        
        # Generate and embed high-resolution Process Flowchart diagram PNG
        flowchart_img = generate_process_flowchart(session_id, sol_name, workflow_steps)
        if flowchart_img:
            self._add_diagram_image(doc, flowchart_img, f"Process Flowchart & Step Execution Architecture for {sol_name}")

        if workflow_steps:
            table_wf = doc.add_table(rows=1, cols=3)
            table_wf.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_wf = table_wf.rows[0].cells
            hdr_wf[0].text = "Step #"
            hdr_wf[1].text = "Workflow Stage"
            hdr_wf[2].text = "Operational & Technical Process"
            for c in hdr_wf:
                _set_cell_background(c, "0F766E")
                _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
                c.paragraphs[0].runs[0].font.bold = True
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

            for idx, step in enumerate(workflow_steps):
                row_wf = table_wf.add_row().cells
                row_wf[0].text = step.get("step", f"Step {idx+1}")
                row_wf[1].text = step.get("title", f"Stage {idx+1}")
                row_wf[2].text = step.get("description", "")
                for c in row_wf:
                    _set_cell_background(c, "F0FDF4" if idx % 2 == 0 else "FFFFFF")
                    _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                    c.paragraphs[0].runs[0].font.size = Pt(8.5)

        # 3. Flowchart: Current vs Target State Comparison
        self._add_heading(doc, "3. Current (As-Is) vs. Target (To-Be) State Matrix", level=1)
        state_info = brd_data.get("current_vs_target_state", {})
        curr_state = state_info.get("current_state", "Manual processing with high operational friction.")
        targ_state = state_info.get("target_state", f"Automated execution powered by {sol_name}.")

        # Generate and embed Current vs Target Comparison diagram PNG
        comp_img = generate_current_vs_target_diagram(session_id, curr_state, targ_state)
        if comp_img:
            self._add_diagram_image(doc, comp_img, "Operational Transformation: Current State vs Target Architecture")

        tbl_st = doc.add_table(rows=1, cols=2)
        tbl_st.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_st = tbl_st.rows[0].cells
        hdr_st[0].text = "❌ Current State (As-Is)"
        hdr_st[1].text = f"✅ Target State ({sol_name})"
        for c in hdr_st:
            _set_cell_background(c, "0F766E")
            _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        row_st = tbl_st.add_row().cells
        row_st[0].text = curr_state
        row_st[1].text = targ_state
        for c in row_st:
            _set_cell_background(c, "F0FDF4")
            _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.size = Pt(9.0)

        # 4. Key Business Benefits & Projected ROI
        self._add_benefits_roi_box(doc, None, context)

        # 5. Scope Boundaries Matrix
        self._add_heading(doc, "5. Solution Scope Boundaries", level=1)
        tbl_scope = doc.add_table(rows=1, cols=2)
        tbl_scope.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_in, c_out = tbl_scope.rows[0].cells
        c_in.text = "In-Scope Capabilities"
        c_out.text = "Out-of-Scope Items"
        _set_cell_background(c_in, "0F766E")
        _set_cell_background(c_out, "0F766E")
        _set_cell_margins(c_in, top=100, bottom=100, left=120, right=120)
        _set_cell_margins(c_out, top=100, bottom=100, left=120, right=120)
        c_in.paragraphs[0].runs[0].font.bold = True
        c_in.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c_out.paragraphs[0].runs[0].font.bold = True
        c_out.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

        row_sc = tbl_scope.add_row().cells
        row_sc[0].text = "\n".join([f"• {item}" for item in brd_data.get("scope_in", ["Core system automation", "Integrations", "Dashboard reporting"])])
        row_sc[1].text = "\n".join([f"• {item}" for item in brd_data.get("scope_out", ["Legacy database overhaul", "Unrelated third-party add-ons"])])
        for c in row_sc:
            _set_cell_background(c, "F0FDF4")
            _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.size = Pt(8.5)

        # 6. Stakeholder Governance Matrix
        self._add_heading(doc, "6. Stakeholder Governance Matrix", level=1)
        table_g = doc.add_table(rows=1, cols=3)
        table_g.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_g = table_g.rows[0].cells
        hdr_g[0].text = "Stakeholder Role"
        hdr_g[1].text = "Key Responsibilities"
        hdr_g[2].text = "Approval Authority"
        for c in hdr_g:
            _set_cell_background(c, "0F766E")
            _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            c.paragraphs[0].runs[0].font.bold = True

        stakeholders = brd_data.get("stakeholders", ["Executive Sponsor", "Digital Transformation Lead", "Consulting Lead", "Technical Architect"])
        for idx, sh in enumerate(stakeholders):
            row_g = table_g.add_row().cells
            row_g[0].text = sh
            row_g[1].text = f"Validate requirement specs, oversee UAT benchmarks, and drive adoption for {sol_name}."
            row_g[2].text = "Executive Sign-off & Budget" if idx == 0 else "Technical / Operational Sign-off"
            for c in row_g:
                _set_cell_background(c, "F0FDF4" if idx % 2 == 0 else "FFFFFF")
                _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                c.paragraphs[0].runs[0].font.size = Pt(8.5)

        file_name = f"BRD_{session_id[:8]}.docx"
        file_path = os.path.join(self.exports_dir, file_name)
        doc.save(file_path)
        logger.info(f"Generated BRD document with embedded visual diagrams at {file_path}")
        return file_path

    def generate_prd(self, session_id: str, prd_data: dict, success_criteria: dict = None, context: dict = None) -> str:
        """Generates visual Project Requirement Document (.docx) with system architecture diagrams and visual tables."""
        doc = Document()
        doc_code = prd_data.get("doc_id_code", f"PRD-2026-{session_id[:6].upper()}")
        sol_name = (context.get("solution_name") if context else None) or prd_data.get("title", "Project Requirement Document")

        self._add_styled_header(doc, f"Project Requirement Document - {sol_name}", "Digitalization Advisor - Technical Architecture Specification", doc_code)

        prod_overview = (context.get("proposed_solution_details") if context else None) or prd_data.get("product_overview", "")
        if prod_overview:
            self._add_green_callout_box(doc, "Product Vision & Architecture Overview", prod_overview)

        # 1. System Architecture Diagram & Technical Stack
        self._add_heading(doc, "1. System Architecture & Technical Layer Diagram", level=1)
        tech_list = (context.get("technologies") if context else None) or prd_data.get("technologies", []) or []
        
        # Generate and embed Technical System Architecture diagram PNG
        arch_img = generate_architecture_diagram(session_id, sol_name, tech_list)
        if arch_img:
            self._add_diagram_image(doc, arch_img, f"System Architecture & Layer Stack Diagram for {sol_name}")

        if tech_list:
            tbl_t = doc.add_table(rows=1, cols=3)
            tbl_t.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_t = tbl_t.rows[0].cells
            hdr_t[0].text = "Technology Component"
            hdr_t[1].text = "Architecture Layer"
            hdr_t[2].text = "Integration Role"
            for c in hdr_t:
                _set_cell_background(c, "0F766E")
                _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
                c.paragraphs[0].runs[0].font.bold = True
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

            for idx, t in enumerate(tech_list):
                row_t = tbl_t.add_row().cells
                row_t[0].text = t
                row_t[1].text = "AI Core Engine / API Layer" if idx == 0 else "Backend & Database Service"
                row_t[2].text = f"Primary service driver for {sol_name} execution."
                for c in row_t:
                    _set_cell_background(c, "F0FDF4" if idx % 2 == 0 else "FFFFFF")
                    _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                    c.paragraphs[0].runs[0].font.size = Pt(8.5)

        # 2. Solution Execution Flowchart
        self._add_heading(doc, "2. Core Process Flowchart & Step Execution", level=1)
        workflow_steps = (context.get("workflow_steps") if context else None) or []
        flowchart_img = generate_process_flowchart(session_id, sol_name, workflow_steps)
        if flowchart_img:
            self._add_diagram_image(doc, flowchart_img, "Solution Execution Flowchart & Task Sequence")

        # 3. System Functional Requirements Matrix
        self._add_heading(doc, "3. System Functional Requirements Matrix", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Module Name"
        hdr_cells[1].text = "Priority"
        hdr_cells[2].text = "Detailed Functional Specification"
        for c in hdr_cells:
            _set_cell_background(c, "0F766E")
            _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            c.paragraphs[0].runs[0].font.bold = True

        funcs = (context.get("functional_requirements") if context else None) or prd_data.get("functional_requirements", []) or []
        for idx, req in enumerate(funcs):
            row_cells = table.add_row().cells
            row_cells[0].text = req.get("module", f"Module {idx+1}")
            row_cells[1].text = req.get("priority", "P0 (Must Have)" if idx < 2 else "P1 (High)")
            row_cells[2].text = req.get("description", "")
            for c in row_cells:
                _set_cell_background(c, "F0FDF4" if idx % 2 == 0 else "FFFFFF")
                _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                c.paragraphs[0].runs[0].font.size = Pt(8.5)

        # 4. Technical & Non-Functional Specifications
        self._add_heading(doc, "4. Technical & Non-Functional Specifications", level=1)
        nfrs = (context.get("technical_requirements") if context else None) or prd_data.get("non_functional_requirements", []) or []
        if nfrs:
            tbl_nfr = doc.add_table(rows=1, cols=2)
            tbl_nfr.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_nfr = tbl_nfr.rows[0].cells
            hdr_nfr[0].text = "Specification Domain"
            hdr_nfr[1].text = "Technical Benchmark Requirement"
            for c in hdr_nfr:
                _set_cell_background(c, "0F766E")
                _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
                c.paragraphs[0].runs[0].font.bold = True
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

            for idx, nfr in enumerate(nfrs):
                row_nfr = tbl_nfr.add_row().cells
                row_nfr[0].text = f"NFR Spec #{idx+1}"
                row_nfr[1].text = nfr
                for c in row_nfr:
                    _set_cell_background(c, "F0FDF4" if idx % 2 == 0 else "FFFFFF")
                    _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                    c.paragraphs[0].runs[0].font.size = Pt(8.5)

        # 5. Success Criteria & Target KPIs
        self._add_benefits_roi_box(doc, success_criteria, context)

        # 6. Risk Assessment & Mitigation Strategy
        self._add_heading(doc, "6. Risk Assessment & Mitigation Matrix", level=1)
        risks_list = (context.get("risks_and_mitigations") if context else None) or []
        if risks_list:
            table_r = doc.add_table(rows=1, cols=2)
            table_r.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_r = table_r.rows[0].cells
            hdr_r[0].text = "Identified Operational Risk"
            hdr_r[1].text = "Target Mitigation Protocol"
            for c in hdr_r:
                _set_cell_background(c, "0F766E")
                _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
                c.paragraphs[0].runs[0].font.bold = True
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

            for idx, r_item in enumerate(risks_list):
                row_r = table_r.add_row().cells
                r_text = r_item.get("risk", "") if isinstance(r_item, dict) else str(r_item)
                m_text = r_item.get("mitigation", "") if isinstance(r_item, dict) else ""
                row_r[0].text = r_text
                row_r[1].text = m_text
                for c in row_r:
                    _set_cell_background(c, "F0FDF4" if idx % 2 == 0 else "FFFFFF")
                    _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                    c.paragraphs[0].runs[0].font.size = Pt(8.5)

        file_name = f"PRD_{session_id[:8]}.docx"
        file_path = os.path.join(self.exports_dir, file_name)
        doc.save(file_path)
        logger.info(f"Generated PRD document with embedded visual diagrams at {file_path}")
        return file_path

    def generate_plan(self, session_id: str, plan_data: dict, context: dict = None) -> str:
        """Generates visual Implementation Plan (.docx) with roadmap charts and visual timelines."""
        doc = Document()
        doc_code = plan_data.get("doc_id_code", f"PLAN-2026-{session_id[:6].upper()}")
        sol_name = context.get("solution_name") if context else plan_data.get("title", "Implementation Plan")

        self._add_styled_header(doc, f"Implementation Plan - {sol_name}", "Digitalization Advisor - Phased Execution Roadmap & Governance", doc_code)

        charter = context.get("project_charter") if context else plan_data.get("project_charter", f"Execution roadmap for {sol_name}.")
        if charter:
            self._add_green_callout_box(doc, "Project Execution Charter & Governance", charter)

        # 1. Phased Solution Delivery Roadmap (Chart Diagram)
        self._add_heading(doc, "1. Phased Delivery Roadmap & Milestone Timeline", level=1)
        phases = (context.get("implementation_phases") if context else None) or plan_data.get("phases", []) or []
        
        # Generate and embed Phased Roadmap Chart diagram PNG
        roadmap_img = generate_roadmap_timeline(session_id, phases)
        if roadmap_img:
            self._add_diagram_image(doc, roadmap_img, f"Phased Implementation Roadmap & Key Deliverables Chart for {sol_name}")

        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Phase & Duration"
        hdr_cells[1].text = f"Implementation Tasks ({sol_name})"
        hdr_cells[2].text = "Tangible Deliverables"
        for c in hdr_cells:
            _set_cell_background(c, "0F766E")
            _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            c.paragraphs[0].runs[0].font.bold = True

        for idx, phase in enumerate(phases):
            row_cells = table.add_row().cells
            row_cells[0].text = f"{phase.get('phase', f'Phase {idx+1}')}\n[{phase.get('duration', '')}]"
            row_cells[1].text = phase.get("tasks", f"Execute development benchmarks for {sol_name}.")
            row_cells[2].text = phase.get("deliverables", "")
            for c in row_cells:
                _set_cell_background(c, "F0FDF4" if idx % 2 == 0 else "FFFFFF")
                _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                c.paragraphs[0].runs[0].font.size = Pt(8.5)

        # 2. System Architecture Reference Diagram
        self._add_heading(doc, "2. Deployment Architecture Reference", level=1)
        arch_img = generate_architecture_diagram(session_id, sol_name, (context.get("technologies") if context else None) or [])
        if arch_img:
            self._add_diagram_image(doc, arch_img, "Deployment & Integration Architecture Map")

        # 3. Key Business Benefits & Projected ROI
        self._add_benefits_roi_box(doc, None, context)

        # 4. Team Resource Allocation & RACI Structure
        self._add_heading(doc, "4. Team Resource Allocation & RACI Structure Matrix", level=1)
        resources = plan_data.get("resource_allocation", ["Lead Solution Architect", "Full-Stack Developer", "Domain Subject Matter Expert"])
        tbl_raci = doc.add_table(rows=1, cols=3)
        tbl_raci.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_raci = tbl_raci.rows[0].cells
        hdr_raci[0].text = "Project Role"
        hdr_raci[1].text = "Core Execution Scope"
        hdr_raci[2].text = "RACI Designation"
        for c in hdr_raci:
            _set_cell_background(c, "0F766E")
            _set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

        for idx, res in enumerate(resources):
            row_r = tbl_raci.add_row().cells
            row_r[0].text = res
            row_r[1].text = f"Lead technical implementation and delivery of {sol_name}."
            row_r[2].text = "Responsible & Accountable" if idx == 0 else "Consulted & Informed"
            for c in row_r:
                _set_cell_background(c, "F0FDF4" if idx % 2 == 0 else "FFFFFF")
                _set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                c.paragraphs[0].runs[0].font.size = Pt(8.5)

        file_name = f"Implementation_Plan_{session_id[:8]}.docx"
        file_path = os.path.join(self.exports_dir, file_name)
        doc.save(file_path)
        logger.info(f"Generated Implementation Plan document with embedded visual diagrams at {file_path}")
        return file_path

doc_generator = DocumentGenerator()

